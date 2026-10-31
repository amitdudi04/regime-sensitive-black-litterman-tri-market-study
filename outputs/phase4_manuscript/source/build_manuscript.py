"""Read frozen publication evidence; create manuscript and cell-level provenance.
No model fitting, return recomputation, data acquisition or optimization occurs here.
Run with the bundled document-runtime Python; export DOCX to PDF using export_word_pdf.ps1.
"""
import argparse, csv, hashlib, json, re
from pathlib import Path
from copy import deepcopy
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ap=argparse.ArgumentParser()
ap.add_argument('--repo',type=Path,default=Path(__file__).resolve().parents[3])
ap.add_argument('--output-dir',type=Path,default=Path(__file__).resolve().parents[1])
args=ap.parse_args(); repo=args.repo.resolve(); out=args.output_dir.resolve(); out.mkdir(parents=True,exist_ok=True)
src=Path(__file__).resolve().parent; frozen=repo/'results/phase3_final'
paths={i:next(frozen.glob(f'table_{i:02d}_*.csv')) for i in range(1,9)}
rows={i:list(csv.DictReader(p.open(encoding='utf-8-sig',newline=''))) for i,p in paths.items()}
refs=json.loads((src/'references.json').read_text(encoding='utf-8')); refmap={r['id']:r for r in refs}
trace=[]; cited=Counter(); table_cells=[]; image_records=[]; section='Front matter'
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def fmt(value, form):
    if value=='': return '—'
    if form=='str': return value
    x=float(value)
    if form.startswith('pct'): return f'{100*x:.{form[3:]}f}%'
    if form.startswith('f'): return f'{x:.{form[1:]}f}'
    if form=='int': return str(round(x))
    if form=='p4': return '<0.0001' if x<.0001 else f'{x:.4f}'
    raise ValueError(form)
def key(row):
    return '; '.join(f'{k}={row[k]}' for k in ['Market','Model','Target Weight Cap','Cost (bps per L1)','Strategy A','Strategy B','Term','Tau','Crisis'] if k in row)
def cell(i,row,col,form,loc):
    display=fmt(row[col],form)
    trace.append(dict(manuscript_location=loc,kind='source_cell',source_path=paths[i].relative_to(repo).as_posix(),source_sha256=sha(paths[i]),source_row_1based=rows[i].index(row)+2,row_key=key(row),source_column=col,source_value=row[col],display_value=display,conversion=form,evidence_phase=row.get('Evidence Phase','')))
    return display
def resolve_token(m):
    bits=m.group(1).split('|'); i=int(bits[0]); market,model,col,form=bits[1:5]
    flags=dict(b.split('=',1) for b in bits[5].split(',')) if len(bits)>5 else {}
    found=[r for r in rows[i] if r.get('Market')==market and (model=='*' or r.get('Strategy B' if i==5 else 'Model')==model)]
    for name,colname in [('cap','Target Weight Cap'),('cost','Cost (bps per L1)'),('tau','Tau'),('term','Term')]:
        if name in flags: found=[r for r in found if (r[colname]==flags[name] if name=='term' else float(r[colname])==float(flags[name]))]
    assert len(found)==1,(bits,len(found))
    return cell(i,found[0],col,form,f'{section}, prose')
def resolve(s):
    s=re.sub(r'\{\{([^}]+)\}\}',resolve_token,s)
    def citation(m):
        ids=m.group(1).split(';')
        for ident in ids: cited[ident]+=1
        return '('+'; '.join(refmap[ident]['cite'] for ident in ids)+')'
    return re.sub(r'\[\[([A-Za-z][A-Za-z0-9;]+)\]\]',citation,s).replace('p=<','p<')

doc=Document()
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(11.5)
normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.12; normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for name,size in [('Title',18),('Heading 1',13),('Heading 2',11.5),('Heading 3',11.5)]:
    sty=doc.styles[name]; sty.font.name='Times New Roman'; sty.font.size=Pt(size); sty.font.color.rgb=RGBColor(0,0,0); sty.font.bold=True
    sty.paragraph_format.space_before=Pt(12); sty.paragraph_format.space_after=Pt(6)
    sty.paragraph_format.keep_with_next=True
    ppr=sty.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in ppr.findall(qn('w:pBdr')): ppr.remove(border)
for sec in doc.sections:
    sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=Inches(.8); sec.left_margin=sec.right_margin=Inches(.8)
    sec.header_distance=sec.footer_distance=Inches(.35)
footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run(); r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
settings=doc.settings.element
update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)

def para(text='',style=None,small=False):
    p=doc.add_paragraph(text,style); p.paragraph_format.widow_control=True
    if small:
        p.paragraph_format.line_spacing=1.05
        for r in p.runs:r.font.size=Pt(9)
    return p
def orientation(landscape):
    # Reuse a newly opened, still empty section for consecutive floats.
    # Otherwise Word creates a blank portrait leaf between landscape pages.
    last=doc.paragraphs[-1]
    fresh=(last.text=='' and last._p.find('./'+qn('w:pPr')+'/'+qn('w:sectPr')) is not None)
    sec=doc.sections[-1] if fresh else doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width=Inches(11 if landscape else 8.5); sec.page_height=Inches(8.5 if landscape else 11)
    sec.top_margin=sec.bottom_margin=Inches(.65 if landscape else .8)
    sec.left_margin=sec.right_margin=Inches(.7 if landscape else .8)
    return sec
def caption(text):
    p=para(text); p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8)
    for r in p.runs:r.bold=True;r.font.size=Pt(10.5)
    return p
def tabulate(i,data,columns,title,notes,land=True):
    if land and not (i==1 and title.startswith('Table 1B')): orientation(True)
    caption(title)
    table=doc.add_table(rows=1,cols=len(columns)); table.autofit=False
    width=(9.6 if land else 6.9)/len(columns)
    for c in table.columns:c.width=Inches(width)
    for j,(label,col,form) in enumerate(columns):table.rows[0].cells[j].text=label
    for n,row in enumerate(data):
        cells=table.add_row().cells
        for j,(label,col,form) in enumerate(columns):
            if callable(col): value=col(row)
            else:value=cell(i,row,col,form,f'{title}; row {n+1}; {label}')
            cells[j].text=value; table_cells.append((len(doc.tables)-1,n+1,j,value))
    header=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
    for n,row in enumerate(table.rows):
        pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            c.width=Inches(width)
            tcpr=c._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'E8E8E8' if n==0 else ('F7F7F7' if n%2==0 else 'FFFFFF'));tcpr.append(shading)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.03
                if i==8:p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j<2 else WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:r.font.size=Pt(9.5 if land else 10);r.bold=(n==0)
    para(notes,small=True)
    if land and not (i==1 and title.startswith('Table 1A')): orientation(False)

short=lambda r:{'Black-Litterman':'BL','Markowitz':'MV','Equal-weight':'1/N'}[r.get('Model',r.get('Strategy B'))]
common='BL = Black-Litterman; MV = historical-mean Markowitz; 1/N = periodically rebalanced equal weighting. Source: frozen Phase 3 publication CSV; underlying Phase 2 experiment. Market-specific OOS: US 2011-01-04–2024-12-31; China 2014-11-07–2024-12-31; India 2013-02-13–2024-12-31. '
def tables(i):
    if i==1:
        tabulate(1,rows[1],[('Market','Market','str'),('ETF universe','ETF Universe','str'),('Currency','Currency','str'),('Common price start','Common Price Start','str'),('Common price end','Common Price End','str'),('Prices N','Common Price Observations','int')],'Table 1A. Dataset: instruments and common adjusted-price histories','Source: Phase 1 dataset provenance, carried into Phase 3 Table 1. Yahoo Finance Adj Close, USD-denominated US-listed ETFs; common observed dates within each basket. Prices are provider-adjusted dividend/split proxies. No missing-price filling.',True)
        tabulate(1,rows[1],[('Market','Market','str'),('OOS start','OOS Start','str'),('OOS end','OOS End','str'),('Returns N','OOS Observations','int'),('Strategy benchmark','Rebalanced Strategy Benchmark','str'),('Passive ETF reference','Passive ETF Reference','str')],'Table 1B. Dataset: rolling OOS samples and benchmark definitions','Source: Phase 3 Table 1. Passive references are portfolio constituents, not independent stand-alone broad-market replications. Later China/India starts reflect observed instrument availability plus the training window.',True)
    if i==2:
        tabulate(2,rows[2],[('Market','Market','str'),('Model',short,'str'),('Net ann. return','Net Annualized Return','pct4'),('Ann. volatility','Annualized Volatility','pct4'),('Sharpe','Sharpe Ratio','f6'),('Max drawdown','Max Drawdown','pct4'),('Mean L1 turnover','Mean Rebalance L1 Turnover','f6'),('Target ASI','ASI','f6'),('Terminal NAV','Net Terminal Wealth','f4')],'Table 2. Principal net performance at 10 bps per L1 turnover',common+'Arithmetic annual return; zero-RF Sharpe. Turnover excludes initial deployment; ASI excludes the first undefined target difference. Initial NAV = 1; initial turnover = 1 is separately charged. Drawdown includes initial capital. No 40% cap.',True)
    if i==3:
        tabulate(3,rows[3],[('Market','Market','str'),('Model',short,'str'),('Target cap','Target Weight Cap','pct2'),('Net ann. return','Net Annualized Return','pct4'),('Sharpe','Sharpe Ratio','f6'),('Mean L1 turnover','Mean Rebalance L1 Turnover','f6'),('Target ASI','ASI','f6'),('Mean target HHI','Mean Target HHI','f6'),('Max target weight','Maximum Target Weight','pct2')],'Table 3. Matched concentration constraints at 10 bps',common+'HHI is the sum of squared target weights. The 40% cap applies only at target formation; holdings drift between trades. BL cap is inactive. Return changes, maximum end-of-day weights and end-of-day HHI remain available in the source CSV.',True)
    if i==4:
        base=[r for r in rows[4] if float(r['Cost (bps per L1)'])==0]
        for panel,col,form,meaning in [('A','Net Annualized Return','pct4','Arithmetic annualized net return'),('B','Net Terminal Wealth','f4','Compounded terminal NAV, initial NAV 1'),('C','Total Costs Per Initial NAV','f6','Direct modeled trading fees in units of initial NAV')]:
            columns=[('Market','Market','str'),('Model',short,'str'),('Target cap','Target Weight Cap','pct2')]
            if panel=='A':columns.append(('Gross ann. return','Gross Annualized Return','pct4'))
            for rate in [0,5,10,20]:
                def get(r,rate=rate,col=col,form=form):
                    match=next(x for x in rows[4] if x['Market']==r['Market'] and x['Model']==r['Model'] and x['Target Weight Cap']==r['Target Weight Cap'] and float(x['Cost (bps per L1)'])==rate)
                    return cell(4,match,col,form,f'Table 4{panel}; {key(match)}')
                columns.append((f'{rate} bps / L1',get,form))
            tabulate(4,base,columns,f'Table 4{panel}. Cost sensitivity: {meaning.lower()}',common+'All 60 source cost scenarios appear in each panel; 15 strategy/market/cap combinations × four rates. Gross paths and targets are held fixed across rates. Panel C sums pretrade net NAV × fee fraction; it is not the gross-minus-net terminal-wealth difference.',True)
    if i==5:
        tabulate(5,rows[5],[('Market','Market','str'),('BL versus',short,'str'),('BL target cap','Target Weight Cap','pct2'),('Sharpe difference','Sharpe Difference A minus B','f6'),('95% CI lower','95% CI Lower','f6'),('95% CI upper','95% CI Upper','f6'),('Two-sided p','Two-sided P-value','p4'),('Paired N','Paired Observations','int')],'Table 5. Paired bootstrap uncertainty for net Sharpe differences',common+'Net 10 bps. Strategy A is BL. MV comparator has the same cap; equal weights already satisfy either cap. Circular paired 20-session blocks; 2,000 draws; seed 42. Percentile intervals and empirical-mean-centered approximate probabilities, unadjusted for multiplicity. Duplicated cap-inactive BL–1/N rows are not independent evidence. All intervals include zero; none of the unadjusted p-values is below 0.05.',True)
    if i==6:
        tabulate(6,rows[6],[('Market','Market','str'),('Model',short,'str'),('Event start','Crisis First Trading Day','str'),('Event end','Crisis Last Trading Day','str'),('Event N','Crisis Observations','int'),('Max drawdown','Buy-Hold Max Drawdown','pct4'),('MDD trough','Buy-Hold Trough Date','str'),('Recovery date','Buy-Hold Recovery Date','str'),('Recovery sessions','Buy-Hold Recovery Trading Days','int')],'Table 6. Standalone historical frozen-holdings stress tests','Source: Phase 1 primary buy-and-hold stress summaries, frozen in Phase 3 Table 6. Gross wealth starts at NAV 1 before the first event return; last pre-event targets are purchased and then held through recovery. MDD uses the running wealth peak including initial capital. Recovery targets crisis-start NAV 1, not the running peak; duration counts trading-index positions from the crisis-window MDD trough. Separate from the main rolling OOS results; no verified daily crisis curve is supplied.',True)
    if i==7:
        tabulate(7,rows[7],[('US model',short,'str'),('Term','Term','str'),('Coefficient','Coefficient','f6'),('HAC t','HAC t-statistic','f4'),('HAC p','HAC p-value','p4'),('Annual alpha','Alpha Annualized Arithmetic','pct4'),('R-squared','R-squared','f6')],'Table 7. US four-factor exposures: separate Phase 1 diagnostic','Source: Phase 1 US net strategy series and Kenneth French daily factors, frozen in Phase 3 Table 7. Sample 2011-01-04–2024-12-31; N = 3,521; HAC maximum lags = 5. Daily portfolio excess returns subtract actual RF; source factor percentages divided by 100. Alpha coefficient is daily; annual alpha = 252 × daily intercept and is shown only on Alpha rows. Other coefficients are dimensionless. No capped or 1/N regressions and no China/India regressions are inferred.',True)
    if i==8:
        tabulate(8,rows[8],[('Fixed tau','Tau','f2'),('US net Sharpe','Sharpe Ratio','f6'),('OOS start','OOS Start','str'),('OOS end','OOS End','str'),('Returns N','OOS Observations','int')],'Table 8. US fixed-tau sensitivity: separate Phase 1 diagnostic','Source: Phase 1 fixed-tau US runs, frozen in Phase 3 Table 8. Net 10 bps per L1, zero-RF Sharpe, equal-weight proxy; Omega = 10 × diagonal annualized covariance, independent of tau. Each row is a fixed setting, not a realization of the adaptive main rule. Only these five tested settings support the interpretation.',False)

figfiles=['01_cumulative_wealth','02_allocation_stability','03_trading_turnover','04_matched_concentration','05_transaction_costs','06_statistical_uncertainty','08_us_factor_exposures','09_us_tau_sensitivity']
figtitles=['Compounded net wealth at 10 bps','Mean target-weight allocation instability','Mean subsequent rebalance L1 turnover','Matched 40% concentration robustness','Compounded terminal wealth across modeled cost rates','Paired bootstrap Sharpe differences and 95% intervals','US four-factor exposure coefficients','US fixed-tau net Sharpe sensitivity']
fignotes=[common+'Net wealth includes initial deployment costs and begins from NAV 1. Arithmetic mean and terminal wealth rankings can differ.',common+'ASI is a mean L1 distance between successive targets, not the amount traded.',common+'Dimensionless weight turnover. Initial cash deployment is excluded from the displayed subsequent rebalance means.',common+'Net 10 bps; target cap applies at trades and may be exceeded by later drift.',common+'Rates are basis points per unit of L1 turnover; return is annualized arithmetic, not CAGR.',common+'Net 10 bps. Percentile paired circular-block intervals, approximate centered probabilities and unadjusted comparisons; all intervals include zero. See Table 5.', 'Source: Phase 1 US diagnostic through frozen Phase 3 Figure 08 and Table 7. US 2011-01-04–2024-12-31, N = 3,521; daily excess-return regression, HAC lag 5. Exposures are conditional associations, not causal mechanisms.', 'Source: Phase 1 fixed-tau US diagnostic through frozen Phase 3 Figure 09 and Table 8. US 2011-01-04–2024-12-31, N = 3,521; net 10 bps, zero-RF Sharpe. Fixed settings differ from the adaptive main rule.']
def figure(i):
    land=i in [1,4,5,6]
    preceding=None
    if i==4 and doc.paragraphs[-1].text.startswith('The cap also reduces'):
        preceding=doc.paragraphs[-1]._p;preceding.getparent().remove(preceding)
    if land:orientation(True)
    if preceding is not None:doc._body._body.insert(-1,preceding)
    caption(f'Figure {i}. {figtitles[i-1]}')
    path=frozen/'figures'/f'{figfiles[i-1]}.png'
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Inches(9.4 if land else 6.8))
    note=fignotes[i-1]
    if i==5:note=common+'Rates are basis points per unit of L1 turnover. The plotted outcome is compounded terminal NAV, not annualized arithmetic return; see Table 4A for arithmetic returns.'
    if i==7:note+=' The figure shows coefficient point estimates; HAC statistics are provided in Table 7.'
    para(note+' The embedded image is the unchanged verified Phase 3 PNG.',small=True)
    record=dict(manuscript_location=f'Figure {i}',kind='figure',source_path=path.relative_to(repo).as_posix(),source_sha256=sha(path),source_row_1based='',row_key='',source_column='',source_value='',display_value=figtitles[i-1],conversion='unchanged PNG; scaled proportionally',evidence_phase='Phase 2 via Phase 3' if i<=6 else 'Phase 1 via Phase 3')
    trace.append(record);image_records.append(record)
    if land:orientation(False)

# Native editable Office mathematics, transformed from MathML using Microsoft's stylesheet.
def mi(x):return '<mi>'+x+'</mi>'
def mo(x):return '<mo>'+x+'</mo>'
def mn(x):return '<mn>'+str(x)+'</mn>'
def mt(x):return '<mtext>'+x+'</mtext>'
def row(*x):return '<mrow>'+''.join(x)+'</mrow>'
def sub(x,y):return '<msub>'+x+y+'</msub>'
def sup(x,y):return '<msup>'+x+y+'</msup>'
def frac(x,y):return '<mfrac>'+x+y+'</mfrac>'
def idx(x,y):return sub(mi(x),mi(y))
def par(x):return row(mo('('),x,mo(')'))
def norm(x):return row(mo('‖'),x,mo('‖'))
sumk='<munderover><mo>∑</mo><mrow><mi>k</mi><mo>=</mo><mn>2</mn></mrow><mi>K</mi></munderover>'
equations={
1:row(idx('r','i,t'),mo('='),frac(idx('P','i,t'),idx('P','i,t−1')),mo('−'),mn(1)),
2:row(sub(mt('max'),mi('w')),par(row(sup(mi('w'),mi('T')),mi('μ'),mo('−'),frac(mi('λ'),mn(2)),sup(mi('w'),mi('T')),mi('Σ'),mi('w'))),mt('  s.t. '),sup(mi('1'),mi('T')),mi('w'),mo('='),mn(1),mo(','),mn(0),mo('≤'),idx('w','i'),mo('≤'),mi('u')),
3:row(mi('Σ'),mo('='),mn(252),par(row(par(row(mn(1),mo('−'),mi('δ'))),mi('S'),mo('+'),mi('δ'),frac(mt('tr(S)'),mi('N')),mi('I')))),
4:row(mi('Π'),mo('='),mi('λ'),mi('Σ'),idx('w','eq'),mo(','),idx('w','eq'),mo('='),frac(mi('1'),mi('N'))),
5:row(mi('P'),mo('='),mi('I'),mo(','),mi('Q'),mo('='),mi('Π'),mo('+'),mn('0.10'),par(row(idx('μ','hist'),mo('−'),mi('Π'))),mo(','),idx('Ω','ii'),mo('='),mt('max'),par(row(mn(10),idx('Σ','ii'),mo(','),sup(mn(10),mn('−12'))))),
6:row(idx('μ','BL'),mo('='),mi('Π'),mo('+'),mi('τ'),mi('Σ'),sup(mi('P'),mi('T')),sup(par(row(mi('Ω'),mo('+'),mi('P'),mi('τ'),mi('Σ'),sup(mi('P'),mi('T')))),mn('−1')),par(row(mi('Q'),mo('−'),mi('P'),mi('Π')))),
7:row(idx('r','p,t'),mo('='),sup(idx('h','t'),mi('T')),idx('r','t'),mo(','),sup(idx('h','i,t+1'),mo('−')),mo('='),frac(row(idx('h','i,t'),par(row(mn(1),mo('+'),idx('r','i,t')))),row(mn(1),mo('+'),idx('r','p,t')))),
8:row(idx('T','k'),mo('='),sub(norm(row(idx('w','k'),mo('−'),sup(idx('h','k'),mo('−')))),mn(1)),mo(','),idx('T','initial'),mo('='),mn(1)),
9:row(mn(1),mo('+'),sup(idx('r','p,t'),mt('net')),mo('='),par(row(mn(1),mo('+'),idx('r','p,t'))),par(row(mn(1),mo('−'),mi('c'),idx('T','k'))),mt('  on trade dates')),
10:row(mt('ASI'),mo('='),frac(mn(1),row(mi('K'),mo('−'),mn(1))),sumk,sub(norm(row(idx('w','k'),mo('−'),idx('w','k−1'))),mn(1))),
11:row(idx('μ','ann'),mo('='),mn(252),mi('r̄'),mo(','),idx('σ','ann'),mo('='),row('<msqrt><mn>252</mn></msqrt>',idx('s','r')),mo(','),mt('SR'),mo('='),frac(idx('μ','ann'),idx('σ','ann')),mt('  (RF = 0)')),
12:row(idx('W','0'),mo('='),mn(1),mo(','),idx('W','t'),mo('='),idx('W','t−1'),par(row(mn(1),mo('+'),sup(idx('r','p,t'),mt('net')))),mo(','),mt('MDD'),mo('='),sub(mt('min'),mi('t')),par(row(frac(idx('W','t'),row(sub(mt('max'),mt('0≤s≤t')),idx('W','s'))),mo('−'),mn(1)))),
13:row(mi('p̂'),mo('='),frac(mn(1),mi('B')),'<munderover><mo>∑</mo><mrow><mi>b</mi><mo>=</mo><mn>1</mn></mrow><mi>B</mi></munderover>',mi('1'),mo('{'),mo('|'),sup(mi('d'),mi('*b')),mo('−'),mi('d̄*'),mo('|'),mo('≥'),mo('|'),idx('d','obs'),mo('|'),mo('}')),
14:row(idx('W','t'),mo('='),sub('<mo>∑</mo>',mi('i')),idx('w','i,0'),frac(idx('P','i,t'),idx('P','i,0')),mo(','),idx('R','rec'),mo('='),mt('pos'),par(idx('t','rec')),mo('−'),mt('pos'),par(idx('t','trough')),mt(' ; recovery target W = 1')),
15:row(idx('r','p,t'),mo('−'),idx('RF','t'),mo('='),mi('α'),mo('+'),idx('β','MKT'),idx('MKT','t'),mo('+'),idx('β','SMB'),idx('SMB','t'),mo('+'),idx('β','HML'),idx('HML','t'),mo('+'),idx('β','MOM'),idx('MOM','t'),mo('+'),idx('ε','t')),
}
xslpath=Path(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL')
transform=etree.XSLT(etree.parse(str(xslpath)))
def equation(i):
    p=para();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(10)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.15),WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.85),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    math=transform(etree.fromstring(('<math xmlns="http://www.w3.org/1998/Math/MathML">'+equations[i]+'</math>').encode())).getroot()
    p._p.append(math);p.add_run(f'\t({i})')

p=doc.add_paragraph('Black–Litterman vs. Markowitz:\nStability, Turnover, and Crisis Dynamics\nAcross Global Markets','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in ['Amit Kumar Dudi','Independent Researcher','amitkumardudi04@gmail.com','First version: March 2026','Revised version: 31 October 2026','Author-review working paper']:
    p=para(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Abstract','Heading 1')
blocks=re.split(r'\n\s*\n',(src/'manuscript.md').read_text(encoding='utf-8').strip())
assert blocks.pop(0)=='ABSTRACT'
para(resolve(blocks.pop(0)))
para('Keywords: Black–Litterman; mean–variance optimization; allocation stability; ETF portfolios; trading turnover; bootstrap inference.',small=True)
para('AI assistance disclosure: ChatGPT and Codex assisted with code inspection, validation support, empirical workflow preparation, research documentation, tables, figures and manuscript editing. The full draft declaration at the end of the paper requires author review and approval before submission.',small=True)
doc.add_page_break()
for block in blocks:
    if block.startswith('#'):
        level=len(block)-len(block.lstrip('#'));text=block[level:].strip();section=text
        doc.add_heading(text,level=min(level,3));continue
    m=re.fullmatch(r'\[\[(TABLE|FIGURE|EQ)(\d+)\]\]',block)
    if m:
        {'TABLE':tables,'FIGURE':figure,'EQ':equation}[m[1]](int(m[2]));continue
    if block=='[[REFERENCES]]':
        for ref in refs:
            p=para(ref['text']+' '+ref['url']);p.paragraph_format.first_line_indent=Inches(-.25);p.paragraph_format.left_indent=Inches(.25)
            p.paragraph_format.keep_together=True
        continue
    para(resolve(block.replace('\n',' ')))

assert set(cited)==set(refmap),(set(refmap)-set(cited),set(cited)-set(refmap))
doc.save(out/'Black_Litterman_Revised_Working_Paper_2026.docx')
with (out/'MANUSCRIPT_RESULT_TRACEABILITY.csv').open('w',encoding='utf-8',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=list(trace[0]));writer.writeheader();writer.writerows(trace)
manifest={'evidence_freeze':'da5d40acfc0e71cac51f0fd8bc288606c6c7aa71','source_table_hashes':{str(p.relative_to(repo)):sha(p) for p in paths.values()},'source_rows':{i:len(r) for i,r in rows.items()},'traceability_records':len(trace),'expected_table_cells':table_cells,'figures':image_records,'references':len(refs),'citation_counts':dict(cited),'native_numbered_equations':15,'logical_tables':8,'physical_table_panels':len(doc.tables),'manuscript_source_sha256':sha(src/'manuscript.md'),'references_sha256':sha(src/'references.json')}
(out/'MANUSCRIPT_BUILD_MANIFEST.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k in ['source_rows','traceability_records','references','native_numbered_equations','logical_tables','physical_table_panels']},indent=2))
