"""Author-review documents and narrow manuscript validation. No empirical reruns."""
import argparse, csv, hashlib, json, re, subprocess, zipfile, unicodedata
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from lxml import etree

repo=Path(__file__).resolve().parents[3]; out=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser()
parser.add_argument('--visual-reviewed-pages',type=int,default=0,help='Explicit operator attestation after inspecting every final rendered page')
options=parser.parse_args()
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
base='da5d40acfc0e71cac51f0fd8bc288606c6c7aa71'
original=Path(r'C:\Users\amitk\Downloads\ssrn-6479462 (7).pdf')
original_hash='38621c631b7645124dacbde12fd18892baee95830eb2d6ff91079d3d86785a07'
assert sha(original)==original_hash
manifest=json.loads((out/'MANUSCRIPT_BUILD_MANIFEST.json').read_text(encoding='utf-8'))
refs=json.loads((out/'source/references.json').read_text(encoding='utf-8'))
docx=out/'Black_Litterman_Revised_Working_Paper_2026.docx';pdf=docx.with_suffix('.pdf')
doc=Document(docx);reader=PdfReader(pdf)
text='\n'.join(p.text for p in doc.paragraphs)
pdftext='\n'.join(p.extract_text() for p in reader.pages)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with zipfile.ZipFile(docx) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    media=[hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')]
assert len(doc.inline_shapes)==8
assert len(xml.findall('.//m:oMath',ns))==15
assert len(doc.tables)==11
for t,r,c,v in manifest['expected_table_cells']:assert doc.tables[t].rows[r].cells[c].text==v,(t,r,c,v)
for figure in manifest['figures']:assert figure['source_sha256'] in media
assert not any(m in text for m in ['{{','[[','Ã','Â','�','p=<'])
assert set(manifest['citation_counts'])=={r['id'] for r in refs}
assert min(len(p.extract_text().strip()) for p in reader.pages)>100
# Check complete frozen input directory against committed Git blob bytes, not a rerun.
frozen_files=subprocess.check_output(['git','ls-tree','-r','--name-only',base,'results/phase3_final'],cwd=repo,text=True).splitlines()
for name in frozen_files:
    expected=subprocess.check_output(['git','show',base+':'+name],cwd=repo)
    assert (repo/name).read_bytes()==expected,name
records=list(csv.DictReader((out/'MANUSCRIPT_RESULT_TRACEABILITY.csv').open(encoding='utf-8',newline='')))
for row in records:
    p=repo/row['source_path'];assert sha(p)==row['source_sha256']
    if row['kind']=='source_cell':
        source=list(csv.DictReader(p.open(encoding='utf-8-sig',newline='')))
        assert source[int(row['source_row_1based'])-2][row['source_column']]==row['source_value']
        # Every empirical value appears in the exported PDF (whitespace/line wrapping ignored).
        assert re.sub(r'\s','',row['display_value']) in re.sub(r'\s','',pdftext),row

# Independently compare complete prose paragraphs from DOCX to PDF after typography normalization.
def canonical(s):
    s=unicodedata.normalize('NFKC',s).replace('\u00ad','')
    return ''.join(ch.lower() for ch in s if ch.isalnum())
paragraph_pages=[]
for i,page in enumerate(reader.pages,1):
    lines=page.extract_text().splitlines()
    # Word's extraction order places the page footer first; omit that field
    # when joining paragraphs split across a physical page boundary.
    if lines and lines[0].strip()==str(i):lines=lines[1:]
    paragraph_pages.append('\n'.join(lines))
cpdf=canonical('\n'.join(paragraph_pages))
compared=0
for p in doc.paragraphs:
    if len(p.text)>50 and not p._p.findall('.//m:oMath',ns):
        assert canonical(p.text) in cpdf,p.text[:120]
        compared+=1

parameters=json.loads((repo/'results/phase2_results/PHASE2_RUN_MANIFEST.json').read_text(encoding='utf-8'))
config_path=repo/'results/phase2_results/PHASE2_RUN_MANIFEST.json'
if not any(r['kind']=='configuration' for r in records):
    for k,v in parameters['configuration'].items():
        records.append(dict(manuscript_location='Abstract / Section 5 configuration',kind='configuration',source_path=config_path.relative_to(repo).as_posix(),source_sha256=sha(config_path),source_row_1based='',row_key='JSON configuration.'+k,source_column=k,source_value=json.dumps(v),display_value=str(v),conversion='parameter definition; see text for units',evidence_phase='Phase 2'))
    for k in ['cost_bps','target_weight_caps','bootstrap','statistical_cost_bps']:
        records.append(dict(manuscript_location='Section 5 / table notes',kind='configuration',source_path=config_path.relative_to(repo).as_posix(),source_sha256=sha(config_path),source_row_1based='',row_key='JSON '+k,source_column=k,source_value=json.dumps(parameters[k]),display_value=json.dumps(parameters[k]),conversion='parameter definition; bootstrap components documented in text',evidence_phase='Phase 2'))
    with (out/'MANUSCRIPT_RESULT_TRACEABILITY.csv').open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=records[0].keys());writer.writeheader();writer.writerows(records)

changes=[
('1–2','Title, abstract, keywords, suggested citation','Retain title, author, affiliation and email; distinguish first and revised dates. Rewrite abstract around actual benchmark/constraint/uncertainty evidence; omit unsupported dominance. Remove inconsistent alternative suggested-citation title.','Title page; abstract','Tables 1–8; original title page'),
('3','Introduction','Retain estimation-error motivation and three-market identity. Narrow market-efficiency/causality assertions; explicitly add 1/N and constraint controls.','1','Tables 2–5; foundational literature'),
('3–5','Literature review 2.1–2.6','Retain mean–variance, BL, instability, shrinkage, costs and regimes scholarship. Clarify scaled-identity LW estimator; add direct benchmark, constraints, HAC and inference references. Distinguish cited methods from implemented methods.','2','References; Phase 2 configuration'),
('6–7','Contributions and allocation-stability novelty','Retain implementation focus; withdraw proprietary innovation claim for mean L1 target distance and unverified ownership contribution.','3; 5.6','Table 2; standard L1 definition'),
('7–8','Dataset; Table 1','Replace uniform 2010–2025 coverage with exact common/OOS dates and USD US-listed universes. Distinguish passive ETF references from rebalanced 1/N.','4; Table 1A/B','table_01_dataset_definitions.csv'),
('7–8','Data suppliers and return definition','Replace Yahoo/Refinitiv and daily-log wording with verified Yahoo Adj Close simple-return convention; no new 2025 observations or download. Document adjustments, alignment, missing-data handling and vendor-vintage limitations.','4; 5.1; 8','Phase 1 dataset manifest; Table 1'),
('8–9','Markowitz objective and covariance','Preserve mean–variance objective and legitimate shrinkage motivation. Supply actual risk aversion, annual scaling, long-only/full-investment conditions and common target cap.','5.1–5.2; Equations 1–3','Phase 2 configuration; actual LedoitWolf estimator'),
('9','Equilibrium prior','Replace capitalization-implied wording with equal-weight equilibrium proxy; missing dated capitalization inputs remain an explicit limitation.','5.3; Equation 4','Phase 1/2 manifests; Table 8'),
('9–10','Views, posterior and regime tau','Retain BL updating framework; give actual P, Q, Omega, blend, floor, tau values and training-only volatility comparison. Do not imply independent analyst information.','5.4; Equations 5–6','Phase 1 validation; Phase 2 configuration'),
('9–10','Rolling accounting and stability','Document drifted holdings, chronological training, initial deployment, multiplicative fees and difference between targets and trades. Replace institutional-grade assertions with defined stylized conventions.','5.5–5.8; Equations 7–12','Phase 1 financial validation; Tables 2–4'),
('10','Hypotheses H1–H4','Replace claims of tested volatility/ASI/SOE mechanisms with explicit research questions. No invented reproduction of original tests.','3; 5.9','Original test artifacts unavailable; Table 5 is a separate Sharpe experiment'),
('11–12','Figure 1 and Table 2','Original figure annotations and tabular performance do not agree. Replace both with linked frozen net paths and principal CSV, with exact windows and metrics; retain mixed market outcomes.','6.1; Table 2; Figure 1','table_02_principal_performance.csv; 01_cumulative_wealth.png'),
('12','Performance mechanism narrative','Remove unsupported claims that US factor or ownership estimates explain China/India outcomes; report market-specific descriptive rankings and near-1/N performance.','6.1; 7','Tables 2, 7'),
('12–14','ASI section 8 and Figure 2','Preserve L1 target-distance concept; replace numerical ASI with frozen values, exclude first undefined difference, remove proprietary label and alleged inferential proof.','5.6; 6.2; Figure 2','Table 2; 02_allocation_stability.png'),
('14–15','Rolling Sharpe section 9 and Figure 3','Do not retain an unverified rolling-Sharpe plot as revised evidence. Report source-backed principal Sharpe, matched constraints and paired uncertainty instead.','6.1; 6.3; 6.5','Tables 2, 3, 5; verified figures'),
('15–16','Cost section 10, Figure 4 and Table 3','Replace 15–30-bps historical presentation with predefined 0/5/10/20-bps source grid. Distinguish L1 from half-L1/notional, direct fees from wealth drag, arithmetic means from products. Withdraw universal cost-neutralization claim.','5.5; 5.8; 6.4; Table 4A–C; Figure 5','All 60 rows of table_04_transaction_cost_sensitivity.csv'),
('16–17','Crisis methodology section 11','Replace weighted-log/first-observation normalization with standalone self-financing frozen holdings and pre-return initial NAV. Separate stress tests from rolling OOS.','5.10; Equations 12, 14','Phase 1 primary buy-and-hold convention; Table 6'),
('17–20','Crisis Figures 5–6 and Table 4','Replace historical summary values with primary frozen buy-and-hold results. Omit crisis curves because verified daily stress paths are unavailable; never reverse-engineer them from summaries.','6.6; Table 6','table_06_crisis_stress_tests.csv'),
('19–20','Recovery definition and comparisons','Define recovery to crisis-start NAV 1 from the crisis-window MDD trough, in trading-index positions. Prior-peak recovery is a different legitimate question; do not call start-NAV recovery the only defensible measure or validate portfolio duration with index dates.','5.10; 6.6; Table 6 notes','Table 6 recovery-target/duration fields'),
('16–20','Universal crisis-protection interpretation','Report worse US BL MDD/recovery, China advantage and mixed India drawdown/recovery; no universal downside claim.','6.6; 7; 11','Table 6'),
('21–23','Factor section 12, Figure 7 and Table 5','Replace all regression estimates with frozen Phase 1 US fit; preserve separate provenance and exact aligned sample, percentage conversion, actual RF, daily alpha and HAC lag 5. No new fit.','5.11; 6.7; Table 7; Figure 7','table_07_us_factor_exposures.csv'),
('22–23','Alleged BL momentum t-sign contradiction','CORRECTION TO PRIOR AUDIT: directly inspected original Table 5 (p.22) and narrative (p.23) both report −0.856. The earlier claim of opposite signs is not confirmed by this PDF. New coefficients replace the old record, but no historical sign contradiction is asserted.','Revision log; replacement estimates in 6.7','Original PDF pages 22–23, visually and textually checked; Table 7'),
('21–23','Beta, momentum and alpha interpretation','Remove statistically-indistinguishable-beta wording without a direct equality test, causal momentum attribution and true-alpha language. Restrict inferences to conditional US estimates and their uncertainty.','6.7; 7','Table 7 coefficient/HAC/alpha fields'),
('23','Statistical section 13','Original p=0.0106 does not meet 1%; p=0.0046 does. Historical ASI/volatility tests are not verified. Revised Sharpe procedure is approximate, paired, circular-block, mean-centered, percentile and unadjusted. No equivalence claim.','5.9; 6.5; Table 5; Figure 6','table_05_statistical_comparisons.csv; missing original test artifacts'),
('24–25','SOE/private section 14, Figure 8 and Table 6','Withdraw empirical ownership section and inconsistent figure/table record. No verified dated stock membership, ownership classifications, returns or test artifacts; retain only historical limitation and future-data requirement.','3; 8; 9; 10','Phase 1 historical-reproduction assessment; original PDF'),
('25–26','Regime section 15, Figure 9 and Table 7','Remove unsupported statistically improved regime outcomes and causal insulation claims. Identify smoothed states as ex-post description, separate from training-only tau and prospective inference. Omit unverified old regime table/figure.','2.4; 5.4; 5.11; 8','Phase 1 regime metadata; no causal design'),
('26–27','Tau section 16, Figure 10 and Table 8','Replace five old Sharpe values with fixed-tau Phase 1 outputs; distinguish adaptive main rule; disclose independent Omega and tested-scope limits.','6.8; Table 8; Figure 8','table_08_tau_sensitivity.csv; 09_us_tau_sensitivity.png'),
('28–29','Synthesis section 17','Retain stabilization/implementation motivation but remove universal performance, ownership, national-efficiency and protection explanations. Discuss near-1/N and matched constraints as alternative explanations.','7; 11','Tables 2–8'),
('29–30','Institutional implications section 18','Replace institutional-grade or guaranteed implementation assertions with stylized trade-offs; no capacity, market-impact or execution-feasibility evidence.','7; 8','Cost accounting assumptions; no microstructure experiment'),
('30','Numerical consistency section 19','Replace plausibility-only assurance with actual CSV-to-table/prose links, original-image fingerprints, manuscript audit and DOCX/PDF substantive comparison.','10; validation report','MANUSCRIPT_RESULT_TRACEABILITY.csv; MANUSCRIPT_AUDIT.json'),
('31–32','Limitations section 20','Retain valid small-universe, ETF aggregation and stylized-cost caveats; expand currency, overlap, selection, unequal samples, prior/view, inference, crisis-path and missing historical-input limits.','8','Tables 1–8 and provenance manifests'),
('32–33','Future section 21','Retain relevant broader-universe/prior/execution questions; label ownership data and prospective regimes as future work; remove unrelated machine-learning expansion.','9','Scope of verified experiments'),
('33–34','Reproducibility section 22','Replace generic full-reproducibility assurances with exact frozen/source commits, runs, pinned Python environment, local snapshot checksum requirements and vendor-access restrictions. Do not claim March replication.','10','Phase 1/2/3 run manifests and requirements lock'),
('34','AI declaration','Replace linguistic-only disclosure with accurate draft covering code inspection, mathematical/test support, workflows, research documentation, tables, figures and editing. Include short front disclosure and full author-approval statement; do not invent completed human review.','Title/abstract page; declarations','Actual Phase 1–4 assistance scope; SSRN AI policy'),
('34–35','Funding, conflicts and availability','Preserve historical no-funding/no-conflict report as prior-version declarations requiring current author confirmation. Do not claim vendor redistribution rights or SSRN acceptance.','Declarations','Original declarations; source-data access limits'),
('35–36','References','Preserve relevant foundations; correct Fernandes issue/pages, Idzorek inspected draft year/spelling, distinguish LW 2004a/b and actual shrinkage source. Add directly relevant benchmark/constraint/HAC/bootstrap sources; list manual metadata review separately.','2; references; reference audit','26-entry checked bibliography; authoritative sources linked'),
('35–36','Omitted references','Remove de Prado ML book and Allen/Powell/Singh tail-risk work because revised text does not use their methods. Replace unverified BL 1990 internal-report metadata with verified 1992 foundational article; no claim that the internal report never existed.','References','Original bibliography compared with actual revised methodological scope'),
]
change_md='# Complete manuscript revision log\n\nRevision: 31 October 2026. Original reviewed from beginning to end (36 pages), including all equations, tables, figures, conclusions and references. Every original section is accounted for below; corrected mathematics and valid scholarship are retained in substance. This is a separately verified revised empirical specification, not exact March numerical replication.\n\n'
change_md+=f'Original located at `{original}`. The requested `(7)(2)` filename was absent; the matching March title/author/36-page `(7)` PDF was inspected. Eight local copies shared the same bytes. Original SHA-256: `{original_hash}`. Original remains unchanged and is not redistributed.\n\nThe Phase 3 register stated PDF unavailable in its earlier environment; that historical statement is preserved. Its availability restriction is resolved for Phase 4. Its unconfirmed momentum-sign allegation is corrected explicitly below.\n\n'
change_md+='| Original PDF pages | Item | Corrected treatment | Revised location | Evidence |\n|---|---|---|---|---|\n'
change_md+='\n'.join('| '+' | '.join(c)+' |' for c in changes)
change_md+='\n\nAll 36 original pages are covered by the locations above. New substantive additions are the equal-weight control, matched cap, complete cost grid, explicit paired uncertainty, clarified accounting/provenance and author-review disclosures; these come from frozen Phases 1–3, not new Phase 4 experiments. Verified figures are renumbered 1–8; source Figure 08 becomes manuscript 7 and source Figure 09 becomes manuscript 8. Table 1 has two panels; Table 4 has three. Old ownership, regime and rolling-Sharpe exhibits are not silently relabeled as verified data.\n'
(out/'MANUSCRIPT_REVISION_CHANGELOG.md').write_text(change_md,encoding='utf-8')

challenges=[
('Does BL add value beyond the chosen prior?','It is nearly identical to rebalanced 1/N on returns and Sharpe. The study establishes stabilization versus unrestricted historical-mean MV, not incremental alpha from Bayesian information.','Tables 2, 5; 6.1–6.2, 7'),
('Is the prior actually a market portfolio?','No. It is an explicitly assumed equal-weight proxy, Pi=lambda Sigma w_eq. Original dated capitalization/AUM inputs remain missing.','5.3, 8; Phase 1/2 configuration'),
('Are the views independent information?','No. Q mildly tilts the implied prior toward historical sample means. Low blend and large diagonal Omega create strong anchoring. No causal decomposition or independent forecast skill is established.','5.4; Equations 5–6'),
('Are units compatible in reverse optimization and posterior?','Annual mean and annual covariance use 252 daily sessions consistently. Realized volatility uses sample daily standard deviation × sqrt(252), distinguished from fitted covariance.','5.1–5.2, 5.6'),
('Does tau cancel from Omega?','Omega is 10 times diagonal annual covariance with a floor, independent of tau. It therefore does not algebraically cancel with tau in this specification. Fixed-tau results are separate US diagnostics.','5.4, 6.8; Table 8'),
('Is the Markowitz comparator artificially weak?','Unrestricted historical means produce strong concentration. Matched 40% caps reduce instability and improve US/China Sharpe while lowering India Sharpe. Both specifications are retained; no claim that the chosen cap is optimal.','Table 3; Figure 4; 6.3'),
('Do the two optimizers receive identical covariance and risk aversion?','Yes, both use the same LW scaled-identity covariance, annual units, lambda=3 and scenario constraints. Inputs differ through the expected-return construction.','5.1–5.4; Phase 2 configuration'),
('What exactly does the 40% cap constrain?','Targets at scheduled trades only. Weights drift afterward and can exceed the target cap. Max end-of-day weights remain separately recorded in the source CSV.','5.7; Table 3 notes'),
('Is equal weighting unfairly rebalanced every day?','No. It resets on the same 63-session schedule and drifts between trades. It pays the same deployment and L1-based modeled fees.','5.5, 5.7; Tables 2, 4'),
('Is ASI the quantity actually traded?','No. ASI is mean successive-target L1 distance excluding the first undefined target difference. Trading turnover compares targets with drifted pretrade holdings. Constant 1/N targets have ASI zero but require trades.','Equations 7–10; Figures 2–3'),
('Are costs funded consistently and counted once?','Fees remove c times L1 turnover from available wealth before the holding-period return. Initial deployment turnover is 1. This is a weight-based proportional model, not an implicit broker-dollar trade-notional equation.','5.5; Equations 8–9; Phase 1 validation'),
('Can fees change gross decision paths in this experiment?','The objective ignores costs and fixed trades use common proportional NAV scaling; drifted weight ratios and gross decisions are unchanged across fee rates. Cost-aware optimization is not claimed.','5.8; Table 4'),
('Why can MV retain higher mean return but lower terminal wealth?','Annual arithmetic mean averages daily returns; terminal wealth multiplies growth factors. US 20-bps MV has higher arithmetic mean and lower terminal NAV than BL. This is not an accounting contradiction.','Table 4A/B; 6.4'),
('Are direct fees the terminal gross–net wealth gap?','No. Fee totals sum pretrade NAV × fee fraction relative to initial NAV. The gross–net terminal gap additionally incorporates foregone investment growth. Neither total is a universal annual fee rate.','Table 4C; 5.6, 6.4'),
('Are national markets directly comparable over a common horizon?','No. USD-listed ETF exposures overlap and have unequal effective dates. The comparisons are within each basket, not identification of local-market efficiency, institutional quality or currency-neutral differences.','Table 1; 4, 8'),
('Does the bootstrap establish significance or equivalence?','No significant Sharpe differences were established. Every interval includes zero; every unadjusted p is >=0.05. This is neither proof of equivalence nor evidence that economically relevant differences are zero.','Table 5; 6.5'),
('Is the inference method an exact Ledoit–Wolf bootstrap?','No. It uses paired 20-session circular blocks, 2,000 draws, seed 42, percentile intervals and approximate empirical-mean centering. It is not studentized and does not guarantee exact finite-sample coverage; one block choice and possible nonstationarity remain limits.','5.9; Equation 13; 8'),
('How are multiple comparisons handled?','They are unadjusted and explicitly labeled. Inactive-cap BL–1/N rows duplicate paths and are not independent evidence. No additional significance is inferred from repeated rows.','Table 5 notes; 5.9, 6.5'),
('Were original volatility and ASI hypothesis tests reproduced?','No. Missing original test artifacts prevent that. Original p=.0106 is not significant at 1%; revised Sharpe tests answer a different question.','Change log p.23; 5.9, 6.5'),
('Does the crisis experiment evaluate adaptive trading through stress?','No. Last pre-event target weights are purchased at event start and held without interim trades. This differs from freezing actual live drifted holdings and from the main rolling simulation.','5.10; Table 6'),
('What does crisis recovery mean?','Recovery targets event-start NAV 1 after the event-window MDD trough, counted by trading-index position difference. It is not previous-peak recovery or elapsed time since event onset. No daily stress curve is fabricated from summaries.','Equation 14; Table 6; 6.6'),
('Does BL provide universal crisis protection?','No. US BL MDD and recovery are worse; China favors BL; India has worse BL MDD and marginally faster start-NAV recovery. Events are selected diagnostics without downside-superiority tests.','Table 6; 6.6'),
('Can the US momentum regression explain China or India?','No. The fit is US-only, conditional and noncausal, with actual RF subtraction and HAC lag 5. No capped or 1/N regressions are invented. Alpha is not statistically positive.','Table 7; 5.11, 6.7'),
('Was an original opposite momentum t-sign actually found?','No. Direct reading of original pp.22–23 gives −0.856 in both places, correcting the earlier documentary audit allegation. New verified regression values differ but do not prove that alleged historical contradiction.','Change log; original PDF hash'),
('Can ex-post regimes support a real-time strategy claim?','No. Smoothed Markov states are retrospective descriptions. The trading tau selector uses only its training-window volatility observations, and no causal or statistically superior regime outcome is asserted.','2.4; 5.4, 5.11; 8'),
('Where is the ownership mechanism evidence?','Original dated stock-level inputs are unavailable. SOE/private empirical claims and old exhibits are withdrawn; obtaining genuine dated inputs is future work.','3, 8–10; change log pp.24–25'),
('Can someone exactly reproduce the March numbers?','Not from the available archive. Frozen revised-study provenance is supplied; original vintages, capitalization priors, ownership data and tests remain missing. Fresh vendor data may differ and raw snapshots are not in Git.','10; Phase 1–3 manifests'),
('Does verified software make the research institutionally implementable?','No. It checks identities and internal consistency, not model optimality, liquidity, capacity, taxes, execution quality or price impact. These remain substantive scope limits.','7–8, 10'),
('What is the contribution if outperformance is insignificant?','A transparent cross-basket comparison of target stability, actual modeled trading, matched constraints, simple anchoring and statistical uncertainty. Mixed outcomes and close 1/N performance sharpen the implementation interpretation.','3, 7, 11'),
('Has the author already approved all scholarship and disclosures?','No. This is an author-review draft. Funding/conflicts, AI use and oversight, author metadata, reference-review items and final interpretations require approval before submission. No SSRN upload occurred.','Declarations; validation report'),
]
review='# Finance reviewer challenge checklist\n\nEvidence-based review for an MSc Finance research portfolio. All answers describe the existing frozen design; no additional experiment is implied.\n\n'
for i,(q,a,e) in enumerate(challenges,1):review+=f'{i}. **{q}**\n\n   {a} Evidence: {e}.\n\n'
(out/'FINANCE_REVIEWER_CHALLENGE_CHECKLIST.md').write_text(review.rstrip()+'\n',encoding='utf-8')

refmd='# Citation and reference audit\n\nChecked on 31 October 2026 against publisher records, primary author copies, Crossref publisher metadata and, for the Grinold/Kahn edition, an authoritative library catalogue. Citation completeness is mechanical; verification is of metadata and the cited methodological scope, not a claim to have independently reproduced each cited study.\n\n'
refmd+='All 26 listed works are cited and all in-text citation keys resolve; no duplicate keys. Ledoit/Wolf 2004a and 2004b are disambiguated. Foundational relevant scholarship is retained; directly relevant additions cover 1/N, constraints, the implemented covariance estimator, inference and HAC.\n\n| Work | Verification finding | Source |\n|---|---|---|\n'
refmd+='\n'.join(f"| {r['id']} | {r['status']} | [Record]({r['url']}) |" for r in refs)
refmd+='''

Manual verification / author decisions before submission:

- Memmel (2003): existence, title, journal and pages are corroborated in the primary Ledoit–Wolf (2008) bibliography. The original Finance Letters article/issue and any original DOI were not independently retrieved. No issue number or DOI is invented; current URL points to the corroborating primary article, not Memmel's original.
- Meucci (2010): inspected author-version date and title support the citation. Optional final Wiley chapter pagination/edition should be checked if a fully specified publisher citation is required. The SSRN DOI identifies the author's evolving record, not a verified Wiley chapter DOI.
- Grinold and Kahn: second-edition copyright year 2000 is supported by the authoritative catalogue; author should confirm the preferred edition/year of the copy actually consulted.
- Idzorek is explicitly the inspected 20 July 2004 draft, rather than silently treating the original 2005 citation as the same dated version. Fernandes (2012) is corrected to 36(5):1354–1361. Hamilton uses DOI 10.2307/1912559.

Removed original references: de Prado (2018) and Allen/Powell/Singh (2012), whose methods are not used in the revised paper. The 1990 BL internal-report citation is replaced by the verified 1992 foundational article; the existence/metadata of the earlier internal report are not adjudicated. Original Markowitz (1959) is cited as the Wiley edition, not a later Yale reprint year.

AI-policy checks used official SSRN guidance: [SSRN AI use](https://www.elsevier.support/ssrn/answer/AI) and [SSRN getting started](https://www.elsevier.support/ssrn/answer/get-started). A short front disclosure and a fuller end declaration are present. Actual human review and final wording require author approval; no completed oversight is invented.
'''
(out/'REFERENCE_VERIFICATION_AND_AUTHOR_REVIEW.md').write_text(refmd,encoding='utf-8')

protected={
'main':'1a82c7599029eb305e60b355b90ac7b80966126f',
'audit/reconcile-crisis-and-tau-pipeline':'77f77d8c6e4fe90d9e81091309fe55d9ce0978ae',
'research/revised-paper-phase1-validation':'598fe0d3c0a7cd7ebaf5709ae62b7c6735b92a0d',
'research/revised-paper-phase2-experiments':'7553cf8bff59a4f3d8bad7d15b779707fc533fa1',
'research/revised-paper-phase3-final-results':base}
for name,expected in protected.items():assert subprocess.check_output(['git','rev-parse','origin/'+name],cwd=repo,text=True).strip()==expected
assert subprocess.check_output(['git','branch','--show-current'],cwd=repo,text=True).strip()=='research/revised-paper-phase4-manuscript'
audit={'status':'PASS_MANUSCRIPT_CHECKS','pdf_pages':len(reader.pages),'logical_tables':8,'table_panels':11,'figures':8,'native_equations':15,'table_cells_checked':len(manifest['expected_table_cells']),'prose_and_caption_paragraphs_matched_docx_to_pdf':compared,'traceability_records':len(records),'references':len(refs),'unresolved_citations':0,'uncited_references':0,'frozen_phase3_files_byte_compared':len(frozen_files),'original_pdf_sha256':sha(original),'docx_sha256':sha(docx),'pdf_sha256':sha(pdf),'protected_branch_baselines':protected,'original_pdf_pages':len(PdfReader(original).pages),'author_approval_complete':False,'numerical_experiments_rerun':False,'full_89_test_suite_rerun':False,'ssrn_submitted':False}
(out/'MANUSCRIPT_AUDIT.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
if options.visual_reviewed_pages:
    assert options.visual_reviewed_pages==len(reader.pages)
    audit['visual_review']={'all_final_pages_inspected':True,'pages':len(reader.pages),'dpi':130,'conversion':'native Microsoft Word PDF via packaged document rasterizer and bundled Poppler','reviewed_pdf_sha256':sha(pdf),'findings':'All pages reviewed at full rendered size; table/figure content and notes visible, native equations legible, numbering continuous, no blank or clipped pages. Separate landscape exhibits intentionally retain whitespace.'}
    (out/'MANUSCRIPT_AUDIT.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
report=f'''# Phase 4 manuscript validation

Status: **READY_FOR_AUTHOR_REVIEW**. This is a complete revised working paper, not an SSRN-approved revision. No SSRN submission, merge, numerical-model edit, historical-results edit or new market-data acquisition was performed.

## Coverage and evidence

- {len(reader.pages)} pages; complete title/abstract, introduction, literature, questions, dataset, methodology, results, discussion, limitations, future work, reproducibility/version history, conclusion, declarations and bibliography.
- Eight logical tables, represented as 11 physical panels (Table 1A/B; Table 4A/B/C). All source rows represented: dataset 3; principal 9; constraints 12; cost 60; inference 12; crisis 6; factor 10; tau 5.
- Eight figures, sequentially numbered; exact source PNG bytes embedded in DOCX. Source 08 → manuscript 7; source 09 → manuscript 8. No crisis curve invented.
- Fifteen editable numbered native Office equations. Return, covariance, objective, prior/posterior, drift, fees, ASI, inference, crisis and factor conventions are explicitly defined.
- {len(records)} provenance records, including empirical table cells, numerical prose, figures and source-configuration definitions. CSV row numbers include the header as row 1; source hashes and raw/display conversions are recorded.
- {len(manifest['expected_table_cells'])} actual saved DOCX table cells equal their expected generated values. Every traced displayed empirical value is present in the PDF. All {len(frozen_files)} frozen Phase 3 files equal the base Git blob bytes; no empirical outputs were overwritten.
- {compared} substantive prose/caption/reference paragraphs matched completely from DOCX into PDF after Unicode/whitespace normalization. All table values separately checked; all 8 image hashes match. PDF was exported directly from the final Word document. Native equation appearance and figure axes/labels require visual review in addition to text checks.

## Interpretation review

Passed: exact unequal OOS dates/US-listed USD instruments; equal-weight proxy and historically derived views; common annual units; chronological 252/63 framework; target ASI versus drift-adjusted L1 trades; initial deployment and single multiplicative fee deduction; matched 40% targets; periodic drifted 1/N; arithmetic means versus wealth and fee totals; approximate bootstrap/percentile/unadjusted inference; all intervals include zero and no p below .05; no exact-equivalence claim; separate gross crisis tests and start-NAV recovery; Phase 1 US-only factors/fixed tau; ex-post regimes versus training-only tau; withdrawn ownership evidence; qualified implementation contribution; distinct historical and revised reproduction claims.

The original PDF was found and all 36 pages inspected. Earlier Phase 3 PDF-unavailable status is historical and is not repeated as a present blocker. One prior audit allegation was corrected: original p.22 table and p.23 narrative both show BL momentum t = −0.856. No original opposite-sign contradiction is asserted. The original file hash still matches `{original_hash}`.

## Citation audit

26 references; zero unresolved citations and zero uncited entries; no duplicate keys. Publisher/primary-author/registry metadata are linked in REFERENCE_VERIFICATION_AND_AUTHOR_REVIEW.md. Corrected Fernandes and Idzorek metadata; separate LW 2004a/b. Memmel original article/issue retrieval remains incomplete; optional final Meucci publisher chapter fields and the author's preferred Grinold/Kahn edition remain manual-review items. Metadata gaps are disclosed rather than invented.

## Formatting and conversion

Word-native DOCX → Microsoft Word ExportAsFixedFormat PDF. The packaged document-skill rasterizer uses that native PDF with bundled Poppler at 130 dpi, since bundled LibreOffice is unavailable on this Windows installation. Native MathML→OMML uses Microsoft's installed MML2OMML stylesheet. No equation images are substituted. Continuous page numbering, portrait prose and landscape wide exhibits, repeated table headers, non-splitting rows, deliberate decimal precision and proportional unchanged figures are used.

Visual QA: see the final completion note below; numerical extraction alone does not certify visual layout. Wide exhibits intentionally use separate landscape pages with whitespace so type remains readable. The script and source text/references, native export wrapper and rendering wrapper are retained under source/. Cross-system pagination may differ with Word/font versions; frozen PDF is the reviewed layout. The DOCX's substantive content is editable.

## Preservation and Git safety

Base `{base}` verified before authoring. Active new branch: `research/revised-paper-phase4-manuscript`. All five protected origin refs equal their recorded baselines in MANUSCRIPT_AUDIT.json; remote refs are checked again after publication. Changes are confined to Phase 4 outputs and their byte-preservation attribute. No force-push, amend or protected ref movement. Historical PDF, paper_results, Phase 1–3 evidence and financial-model code are unchanged. Final branch commit is reported in the delivery message; no self-referential commit hash is embedded in its own contents.

## Remaining research limitations

Small overlapping ETF baskets, USD proxy exposures and unequal dates; vendor adjustment/vintage and universe-selection limits; assumed proxy prior, historical tilt and uncertainty settings; stylized costs without execution/impact/capacity evidence; finite-sample approximate inference and unadjusted multiplicity; selected gross stress events and absent archived daily crisis curves; US-only factor fit; restricted fixed-tau range; missing original March vintages, capitalization, ownership and hypothesis-test artifacts. These limit interpretation and historical replication, while the narrower revised-study evidence remains traceable.

## Author approval before submission

1. Review and accept all revised methods, mixed results, interpretations, withdrawals and reference changes.
2. Confirm author contact/affiliation and revision date.
3. Confirm current funding and conflict-of-interest declarations; historical statements are not assumed current.
4. Approve the accurate AI-assistance/oversight declaration and actual human-review account, including front disclosure and applicable institutional/SSRN requirements.
5. Resolve or knowingly retain the explicitly qualified manual bibliography items; confirm scholarly sources actually consulted.
6. Confirm lawful access/distribution arrangements for frozen vendor inputs and the intended public availability wording.

No task blocker remains for creating this author-review package. These decisions remain outstanding for publication; author approval has not been fabricated. No claim of official SSRN approval or exact March numerical reproduction is made.
'''
(out/'PHASE4_MANUSCRIPT_VALIDATION.md').write_text(report,encoding='utf-8')
if options.visual_reviewed_pages:
    with (out/'PHASE4_MANUSCRIPT_VALIDATION.md').open('a',encoding='utf-8') as f:
        f.write(f'\n## Final visual QA completion\n\nAll {len(reader.pages)} pages of the final PDF were rendered at 130 dpi and visually inspected, including every table/figure, caption, source note, equation, footer and reference page. No blank, clipped or unreadable exhibit remains. Table 1 panels share a landscape page; the Table 8 note stays with its table; wide plots preserve original image proportions. Final PDF SHA-256: `{sha(pdf)}`. The two manuscript formats contain the same substantive paper.\n')
(out/'source/README.md').write_text('''# Manuscript source

Run `build_manuscript.py` using Python with python-docx and lxml, pointing --repo at the frozen research repository if moved. This reads the eight Phase 3 CSV files and eight original PNG figures and creates DOCX, cell traceability and build manifest. It performs no numerical research experiment.

On the recorded Windows environment, run export_word_pdf.ps1 with absolute -Docx and -Pdf paths using Microsoft Word COM; this updates page fields and exports the matching PDF. The installed Microsoft MML2OMML.XSL supports native equations. Then run render_via_word_pdf.py DOCX PDF scratch-render-directory using the bundled document runtime. Its recorded paths identify this environment; update them when moving installations. All rendered page PNGs stay in work/ as QA intermediates, not delivered research outputs.

Run revision_review.py after final export to regenerate the revision log, reviewer checklist, reference audit and narrow manuscript checks. It checks the recorded original PDF path and hash and protected origin refs. Update only the path if relocating the same original bytes. It does not refit or download anything. Visual QA must also be completed; the final validation note records that separate step.

The build is reproducible in content, not guaranteed byte-identical across Word/font/library versions; PDF metadata and pagination can change. Read MANUSCRIPT_BUILD_MANIFEST.json and MANUSCRIPT_AUDIT.json for the frozen reviewed outputs and provenance.
''',encoding='utf-8')
print(json.dumps(audit,indent=2))
